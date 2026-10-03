#!/usr/bin/env python3
"""INDRA Excel Spreadsheet Generator: Context-Agnostic, 100% Offline openpyxl Engine.

Supports:
1. CSV files (.csv)
2. Pipe-delimited or tab-delimited text (.txt)
3. JSON array of objects or table spec (.json)
"""
from __future__ import annotations

import argparse
import csv
import json
import os
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
ZEBRA_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
REGULAR_FONT = Font(name="Calibri", size=11, color="1E293B")
BOLD_FONT = Font(name="Calibri", size=11, bold=True, color="0F172A")
THIN_BORDER = Border(
    left=Side(style="thin", color="E2E8F0"),
    right=Side(style="thin", color="E2E8F0"),
    top=Side(style="thin", color="E2E8F0"),
    bottom=Side(style="thin", color="E2E8F0"),
)

def _is_number(val: str) -> tuple[bool, float | int | str]:
    val_clean = str(val).strip().replace(",", "")
    try:
        if "." in val_clean:
            return True, float(val_clean)
        return True, int(val_clean)
    except ValueError:
        return False, val

def render_xlsx(input_path: str, output_path: str, sheet_name: str = "Data", auto_total: bool = True) -> bool:
    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name

    # Load rows
    content_raw = Path(input_path).read_text(encoding="utf-8-sig")
    rows: list[list[str]] = []

    if input_path.endswith(".json"):
        spec = json.loads(content_raw)
        if isinstance(spec, list) and spec and isinstance(spec[0], dict):
            # Array of dicts
            headers = list(spec[0].keys())
            rows.append(headers)
            for item in spec:
                rows.append([str(item.get(h, "")) for h in headers])
        elif isinstance(spec, dict) and "rows" in spec:
            rows = spec["rows"]
    else:
        # Check delimiter: comma, pipe, or tab
        sample = content_raw[:1024]
        delimiter = "|" if "|" in sample else ("\t" if "\t" in sample else ",")
        reader = csv.reader(content_raw.splitlines(), delimiter=delimiter)
        for r in reader:
            cleaned = [c.strip() for c in r if c.strip() or delimiter != "|"]
            if cleaned:
                rows.append(cleaned)

    if not rows:
        print(f"[indra-xlsx] Warning: No rows found in {input_path}", file=sys.stderr)
        rows = [["Item", "Value"], ["No data", "0"]]

    # Write data
    numeric_cols: dict[int, list[float | int]] = {}
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, val in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx)
            if r_idx == 1:
                # Header row
                cell.value = str(val)
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                is_num, parsed = _is_number(val)
                if is_num:
                    cell.value = parsed
                    numeric_cols.setdefault(c_idx, []).append(parsed)
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                else:
                    cell.value = str(val)
                    cell.alignment = Alignment(horizontal="left", vertical="center")

                cell.font = REGULAR_FONT
                if r_idx % 2 == 0:
                    cell.fill = ZEBRA_FILL
            cell.border = THIN_BORDER

    # Freeze header row
    ws.freeze_panes = "A2"

    # Add Totals row if requested and numeric columns exist
    last_row = len(rows)
    if auto_total and numeric_cols and last_row > 2:
        tot_row = last_row + 1
        ws.cell(row=tot_row, column=1, value="Total").font = BOLD_FONT
        for col_idx in numeric_cols:
            col_letter = get_column_letter(col_idx)
            cell = ws.cell(row=tot_row, column=col_idx)
            cell.value = f"=SUM({col_letter}2:{col_letter}{last_row})"
            cell.font = BOLD_FONT
            cell.alignment = Alignment(horizontal="right", vertical="center")
            cell.border = Border(top=Side(style="thin", color="0F172A"), bottom=Side(style="double", color="0F172A"))

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(str(out))
    print(f"[indra-xlsx] Successfully created spreadsheet: {output_path}")
    return True

def main():
    parser = argparse.ArgumentParser(description="INDRA 100% Offline Excel Spreadsheet Generator")
    parser.add_argument("input", help="Path to input CSV (.csv), text (.txt), or JSON (.json)")
    parser.add_argument("-o", "--output", required=True, help="Destination .xlsx path")
    parser.add_argument("--sheet", default="Data", help="Sheet name")
    args = parser.parse_args()

    render_xlsx(args.input, args.output, args.sheet)

if __name__ == "__main__":
    main()
