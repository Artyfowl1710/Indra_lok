"""Offline document-generation smoke test for the INDRA execution image."""

from pathlib import Path
import sys

from openpyxl import Workbook, load_workbook
from PIL import Image, ImageDraw
from pptx import Presentation
from pypdf import PdfReader
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas


out = Path(sys.argv[1] if len(sys.argv) > 1 else "/workspace/indra-document-smoke")
out.mkdir(parents=True, exist_ok=True)

sample = out / "source.png"
image = Image.new("RGB", (800, 320), "white")
draw = ImageDraw.Draw(image)
draw.text((30, 35), "INDRA demo source: incident report", fill="black")
image.save(sample)

pdf = out / "report.pdf"
page = canvas.Canvas(str(pdf), pagesize=A4)
page.setFont("Helvetica-Bold", 18)
page.drawString(50, 785, "Incident Summary")
page.setFont("Helvetica", 11)
for index in range(12):
    page.drawString(50, 740 - index * 32, f"Finding {index + 1}: verified offline document output")
page.drawImage(str(sample), 50, 125, width=340, height=136)
page.save()
assert len(PdfReader(str(pdf)).pages) == 1

xlsx = out / "findings.xlsx"
book = Workbook()
sheet = book.active
sheet.title = "Findings"
sheet.append(["ID", "Finding", "Severity"])
sheet.append([1, "Offline document output", "Medium"])
book.save(xlsx)
assert load_workbook(xlsx, read_only=True).active["B2"].value == "Offline document output"

pptx = out / "briefing.pptx"
slides = Presentation()
slide = slides.slides.add_slide(slides.slide_layouts[5])
slide.shapes.title.text = "Incident Briefing"
slide.shapes.add_picture(str(sample), 1000000, 1700000, width=6500000)
slides.save(pptx)
assert len(Presentation(pptx).slides) == 1

for path in (pdf, xlsx, pptx):
    print(f"OK {path} ({path.stat().st_size} bytes)")
