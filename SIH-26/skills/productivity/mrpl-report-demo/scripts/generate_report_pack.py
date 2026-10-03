#!/usr/bin/env python3
"""Generate validated MRPL DOCX, PDF, PPTX, and XLSX files from analysis JSON."""
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path


BLUE = "123A63"
TEAL = "008C95"
LIGHT = "EAF2F5"


def _analysis_from_text(path: Path) -> dict:
    """Create a compact, evidence-only report structure from a source text file."""
    text = path.read_text(encoding="utf-8", errors="replace").strip()
    if not text:
        raise ValueError(f"source file is empty: {path}")

    lines = [line.strip() for line in text.splitlines() if line.strip()]
    paragraphs = [part.strip() for part in re.split(r"\n\s*\n", text) if part.strip()]
    summary_parts = paragraphs[:3] if len(paragraphs) > 1 else lines[:5]
    summary = " ".join(summary_parts)
    if len(summary) > 1200:
        summary = summary[:1197].rsplit(" ", 1)[0] + "..."

    sections = []
    current_heading = "Source Findings"
    current_lines: list[str] = []
    for line in lines:
        looks_like_heading = (
            len(line) <= 90
            and not line.endswith((".", ",", ";"))
            and (line.isupper() or line.endswith(":"))
        )
        if looks_like_heading and current_lines:
            sections.append({"heading": current_heading, "paragraphs": current_lines[:6], "bullets": []})
            current_heading = line.rstrip(":")
            current_lines = []
        elif looks_like_heading:
            current_heading = line.rstrip(":")
        else:
            current_lines.append(line)
    if current_lines:
        sections.append({"heading": current_heading, "paragraphs": current_lines[:10], "bullets": []})
    if not sections:
        sections = [{"heading": "Source Findings", "paragraphs": lines[:12], "bullets": []}]

    metric_pattern = re.compile(r"(?:₹|Rs\.?|INR|%|crores?|million|billion|MT|MMTPA)", re.IGNORECASE)
    metrics = [
        {"metric": f"Reported figure {index}", "value": line, "note": "As stated in supplied source"}
        for index, line in enumerate((line for line in lines if metric_pattern.search(line)), 1)
    ][:12]
    return {
        "title": "MRPL Financial Analysis",
        "subtitle": "Analysis of supplied material",
        "executive_summary": summary,
        "sections": sections[:8],
        "metrics": metrics,
        "sources": [path.name],
    }


def load_analysis(path: Path) -> dict:
    if path.suffix.lower() != ".json":
        return _analysis_from_text(path)
    data = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, dict) or not str(data.get("title", "")).strip():
        raise ValueError("analysis JSON requires a non-empty title")
    data.setdefault("subtitle", "Analysis of supplied material")
    data.setdefault("executive_summary", "Not provided in source.")
    data.setdefault("sections", [])
    data.setdefault("metrics", [])
    data.setdefault("sources", [])
    return data


def create_docx(data: dict, path: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    title = doc.add_heading(data["title"], 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph(data["subtitle"])
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph(data["executive_summary"])
    for item in data["sections"]:
        doc.add_heading(str(item.get("heading", "Finding")), level=1)
        for paragraph in item.get("paragraphs", []):
            doc.add_paragraph(str(paragraph))
        for bullet in item.get("bullets", []):
            doc.add_paragraph(str(bullet), style="List Bullet")
    if data["metrics"]:
        doc.add_heading("Key Metrics", level=1)
        table = doc.add_table(rows=1, cols=3)
        table.style = "Light Shading Accent 1"
        for cell, label in zip(table.rows[0].cells, ("Metric", "Value", "Context")):
            cell.text = label
        for metric in data["metrics"]:
            cells = table.add_row().cells
            cells[0].text = str(metric.get("metric", ""))
            cells[1].text = str(metric.get("value", ""))
            cells[2].text = str(metric.get("note", ""))
    if data["sources"]:
        doc.add_heading("Sources", level=1)
        for source in data["sources"]:
            doc.add_paragraph(str(source), style="List Bullet")
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.color.rgb = RGBColor.from_string(BLUE)
    doc.save(path)
    Document(path)


def create_pdf(data: dict, path: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
    from pypdf import PdfReader

    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="MRPLTitle", parent=styles["Title"], textColor=colors.HexColor(f"#{BLUE}")))
    story = [Paragraph(data["title"], styles["MRPLTitle"]), Paragraph(data["subtitle"], styles["Italic"]), Spacer(1, 8)]
    story += [Paragraph("Executive Summary", styles["Heading1"]), Paragraph(data["executive_summary"], styles["BodyText"])]
    for item in data["sections"]:
        story.append(Paragraph(str(item.get("heading", "Finding")), styles["Heading1"]))
        for paragraph in item.get("paragraphs", []):
            story += [Paragraph(str(paragraph), styles["BodyText"]), Spacer(1, 4)]
        for bullet in item.get("bullets", []):
            story.append(Paragraph(f"• {bullet}", styles["BodyText"]))
    if data["metrics"]:
        story.append(Paragraph("Key Metrics", styles["Heading1"]))
        rows = [["Metric", "Value", "Context"]] + [[str(m.get("metric", "")), str(m.get("value", "")), str(m.get("note", ""))] for m in data["metrics"]]
        table = Table(rows, colWidths=[45 * mm, 40 * mm, 85 * mm], repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{BLUE}")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("BACKGROUND", (0, 1), (-1, -1), colors.HexColor(f"#{LIGHT}")),
        ]))
        story.append(table)
    if data["sources"]:
        story.append(Paragraph("Sources", styles["Heading1"]))
        for source in data["sources"]:
            story.append(Paragraph(f"• {source}", styles["BodyText"]))
    SimpleDocTemplate(str(path), pagesize=A4, rightMargin=18 * mm, leftMargin=18 * mm, topMargin=18 * mm, bottomMargin=18 * mm).build(story)
    if not PdfReader(str(path)).pages:
        raise ValueError("generated PDF has no pages")


def create_pptx(data: dict, path: Path) -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = data["title"]
    slide.placeholders[1].text = data["subtitle"]
    summary = prs.slides.add_slide(prs.slide_layouts[1])
    summary.shapes.title.text = "Executive Summary"
    summary.placeholders[1].text = data["executive_summary"]
    for item in data["sections"]:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = str(item.get("heading", "Finding"))
        lines = [str(p) for p in item.get("paragraphs", [])] + [str(b) for b in item.get("bullets", [])]
        frame = slide.placeholders[1].text_frame
        frame.clear()
        for index, line in enumerate(lines[:7]):
            para = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            para.text = line
            para.level = 0
            para.font.size = Pt(20)
    if data["metrics"]:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = "Key Metrics"
        rows, cols = len(data["metrics"]) + 1, 3
        table = slide.shapes.add_table(rows, cols, Inches(0.7), Inches(1.6), Inches(11.9), Inches(4.8)).table
        for col, label in enumerate(("Metric", "Value", "Context")):
            table.cell(0, col).text = label
        for row, metric in enumerate(data["metrics"], 1):
            table.cell(row, 0).text = str(metric.get("metric", ""))
            table.cell(row, 1).text = str(metric.get("value", ""))
            table.cell(row, 2).text = str(metric.get("note", ""))
        for cell in table.rows[0].cells:
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(BLUE)
            for run in cell.text_frame.paragraphs[0].runs:
                run.font.color.rgb = RGBColor(255, 255, 255)
    prs.save(path)
    Presentation(path)


def create_xlsx(data: dict, path: Path) -> None:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary.append([data["title"]])
    summary.append([data["subtitle"]])
    summary.append([])
    summary.append(["Executive Summary"])
    summary.append([data["executive_summary"]])
    summary.column_dimensions["A"].width = 110
    summary["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    summary["A1"].fill = PatternFill("solid", fgColor=BLUE)
    summary["A4"].font = Font(bold=True, color="FFFFFF")
    summary["A4"].fill = PatternFill("solid", fgColor=TEAL)
    summary["A5"].alignment = Alignment(wrap_text=True, vertical="top")
    findings = wb.create_sheet("Findings")
    findings.append(["Section", "Type", "Content"])
    for item in data["sections"]:
        for paragraph in item.get("paragraphs", []):
            findings.append([item.get("heading", ""), "Paragraph", paragraph])
        for bullet in item.get("bullets", []):
            findings.append([item.get("heading", ""), "Finding", bullet])
    metrics = wb.create_sheet("Metrics")
    metrics.append(["Metric", "Value", "Context"])
    for metric in data["metrics"]:
        metrics.append([metric.get("metric", ""), metric.get("value", ""), metric.get("note", "")])
    sources = wb.create_sheet("Sources")
    sources.append(["Supplied source"])
    for source in data["sources"]:
        sources.append([source])
    for sheet in wb.worksheets:
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=BLUE)
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(80, max(14, max(len(str(c.value or "")) for c in column) + 2))
            for cell in column:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)
    load_workbook(path)


CREATORS = {"docx": create_docx, "pdf": create_pdf, "pptx": create_pptx, "xlsx": create_xlsx}


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path, help="Analysis JSON or source TXT/Markdown file")
    parser.add_argument("--outdir", type=Path, default=Path("outputs/mrpl-demo-output"))
    parser.add_argument("--formats", default="docx,pdf")
    args = parser.parse_args()
    data = load_analysis(args.source)
    args.outdir.mkdir(parents=True, exist_ok=True)
    stem = "MRPL_Analysis_Report"
    results = []
    for fmt in [item.strip().lower() for item in args.formats.split(",") if item.strip()]:
        if fmt not in CREATORS:
            raise ValueError(f"unsupported format: {fmt}")
        output = args.outdir / f"{stem}.{fmt}"
        CREATORS[fmt](data, output)
        if not output.is_file() or output.stat().st_size == 0:
            raise ValueError(f"empty output: {output}")
        results.append({"format": fmt, "status": "created", "path": str(output), "bytes": output.stat().st_size})
    print(json.dumps({"status": "ok", "outputs": results}, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
